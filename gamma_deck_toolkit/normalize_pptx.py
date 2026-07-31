#!/usr/bin/env python3
"""Post-export normalizer for Aviation Synergy Gamma decks.

Takes a raw Gamma .pptx export and enforces the measured house tokens that
Gamma's auto-fit will not hold on its own:

  1. Font floor        -- lift any run below MIN_FONT_PT
  2. Colour unification-- fold near-identical navy strays into the canonical navy
  3. Logo placement    -- snap the footer logo to its exact size and position
  4. Margin snapping   -- pull drifted shapes back onto the layout grid
  5. Title audit       -- report titles Gamma shrank (needs a text edit, not a
                          geometry one, so it is reported rather than forced)

Usage:
    python normalize_pptx.py deck.pptx                    # -> deck-normalized.pptx
    python normalize_pptx.py deck.pptx -o out.pptx
    python normalize_pptx.py deck.pptx --dry-run
    python normalize_pptx.py deck.pptx --skip snap,logo
"""

import argparse
import os
import sys
from collections import Counter

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
except ImportError:
    sys.exit("python-pptx is required:  pip install python-pptx")

import avs_tokens as T

ALL_STEPS = ("floor", "colour", "logo", "snap", "titles")


def iter_shapes(container):
    """Walk shapes, descending into groups."""
    for shape in container.shapes:
        if shape.shape_type == 6:  # GROUP
            yield from iter_shapes(shape)
        else:
            yield shape


def iter_runs(shape):
    if not shape.has_text_frame:
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            yield run


def run_rgb(run):
    """Return the run's RGB hex, or None when it is not a plain RGB colour."""
    try:
        colour = run.font.color
        if colour is None or colour.type is None:
            return None
        return str(colour.rgb)
    except (AttributeError, TypeError):
        return None


def max_run_pt(shape):
    sizes = [r.font.size.pt for r in iter_runs(shape) if r.font.size is not None]
    return max(sizes) if sizes else None


def is_section_divider(slide):
    """Section dividers carry a tall portrait image hard against the right edge."""
    for shape in iter_shapes(slide):
        if shape.shape_type != 13:  # PICTURE
            continue
        if (T.inches(shape.left) > 7.0
                and T.inches(shape.height) > 2.5
                and T.inches(shape.width) < 3.0):
            return True
    return False


def looks_like_logo(shape):
    return (
        shape.shape_type == 13
        and T.inches(shape.width) <= T.LOGO["detect_max_width_in"]
        and T.inches(shape.top) >= T.LOGO["detect_min_top_in"]
    )


# --- Steps ------------------------------------------------------------------

def step_floor(slide, slide_no, log):
    changed = 0
    for shape in iter_shapes(slide):
        for run in iter_runs(shape):
            size = run.font.size
            if size is not None and size.pt < T.MIN_FONT_PT:
                log.append(f"s{slide_no}: font {size.pt:g}pt -> {T.MIN_FONT_PT:g}pt "
                           f"| {run.text[:38]!r}")
                run.font.size = Pt(T.MIN_FONT_PT)
                changed += 1
    return changed


def step_colour(slide, slide_no, log):
    changed = 0
    variants = {v.upper() for v in T.NAVY_VARIANTS}
    for shape in iter_shapes(slide):
        for run in iter_runs(shape):
            hexval = run_rgb(run)
            if hexval and hexval.upper() in variants:
                log.append(f"s{slide_no}: #{hexval} -> #{T.NAVY} "
                           f"| {run.text[:38]!r}")
                run.font.color.rgb = RGBColor.from_string(T.NAVY)
                changed += 1
    return changed


def step_logo(slide, slide_no, log):
    changed = 0
    target_x = (T.LOGO["x_section_in"] if is_section_divider(slide)
                else T.LOGO["x_content_in"])
    for shape in iter_shapes(slide):
        if not looks_like_logo(shape):
            continue
        before = (T.inches(shape.left), T.inches(shape.top),
                  T.inches(shape.width), T.inches(shape.height))
        after = (target_x, T.LOGO["top_in"],
                 T.LOGO["width_in"], T.LOGO["height_in"])
        if any(abs(b - a) > 0.005 for b, a in zip(before, after)):
            log.append(
                f"s{slide_no}: logo "
                f"[{before[0]:.2f},{before[1]:.2f} {before[2]:.2f}x{before[3]:.2f}] -> "
                f"[{after[0]:.2f},{after[1]:.2f} {after[2]:.2f}x{after[3]:.2f}]")
            shape.left = T.emu(after[0])
            shape.top = T.emu(after[1])
            shape.width = T.emu(after[2])
            shape.height = T.emu(after[3])
            changed += 1
    return changed


def step_snap(slide, slide_no, log):
    """Pull drifted left edges back onto the grid, but only small drifts --
    anything further out was a deliberate placement."""
    changed = 0
    for shape in iter_shapes(slide):
        if looks_like_logo(shape):
            continue
        left_in = T.inches(shape.left)
        nearest = min(T.COLUMNS, key=lambda c: abs(c - left_in))
        delta = abs(nearest - left_in)
        if T.SNAP_MIN_IN <= delta <= T.SNAP_TOLERANCE_IN:
            log.append(f"s{slide_no}: left {left_in:.3f}in -> {nearest:.2f}in "
                       f"(drift {delta:.3f})")
            shape.left = T.emu(nearest)
            changed += 1
    return changed


def step_titles(slide, slide_no, log):
    """Report-only: a shrunk title needs shorter text, not a bigger box."""
    target = T.TYPE_RAMP["title"]["pt"]
    limit = T.TYPE_RAMP["title"]["max_chars"]
    flagged = 0
    for shape in iter_shapes(slide):
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        size = max_run_pt(shape)
        if size is None or size < 28.0:
            continue
        text = shape.text_frame.text.strip().replace("\n", " ")
        if size < target:
            log.append(
                f"s{slide_no}: title at {size:g}pt (want {target:g}pt) -- "
                f"{len(text)} chars, trim to <={limit}: {text[:46]!r}")
            flagged += 1
        break
    return flagged


STEP_FUNCS = {
    "floor": step_floor,
    "colour": step_colour,
    "logo": step_logo,
    "snap": step_snap,
    "titles": step_titles,
}


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", help="Gamma .pptx export to normalize")
    ap.add_argument("-o", "--output", help="Output path (default: <name>-normalized.pptx)")
    ap.add_argument("--dry-run", action="store_true", help="Report without writing")
    ap.add_argument("--skip", default="",
                    help=f"Comma-separated steps to skip from: {','.join(ALL_STEPS)}")
    ap.add_argument("--quiet", action="store_true", help="Summary only")
    args = ap.parse_args()

    skip = {s.strip() for s in args.skip.split(",") if s.strip()}
    unknown = skip - set(ALL_STEPS)
    if unknown:
        sys.exit(f"unknown step(s): {', '.join(sorted(unknown))}")
    steps = [s for s in ALL_STEPS if s not in skip]

    prs = Presentation(args.pptx)

    w, h = T.inches(prs.slide_width), T.inches(prs.slide_height)
    if abs(w - CANVAS_W) > 0.02 or abs(h - CANVAS_H) > 0.02:
        print(f"WARNING: canvas is {w:.2f}x{h:.2f}in, tokens expect "
              f"{CANVAS_W:.2f}x{CANVAS_H:.2f}in ({T.CANVAS['name']}). "
              f"Grid snapping may be wrong.\n")

    logs = {s: [] for s in steps}
    counts = Counter()

    for n, slide in enumerate(prs.slides, start=1):
        for step in steps:
            counts[step] += STEP_FUNCS[step](slide, n, logs[step])

    if not args.quiet:
        for step in steps:
            entries = logs[step]
            if not entries:
                continue
            verb = "flagged" if step == "titles" else "changed"
            print(f"{step.upper()} -- {len(entries)} {verb}")
            for line in entries:
                print(f"  {line}")
            print()

    print(f"{os.path.basename(args.pptx)}: {len(prs.slides._sldIdLst)} slides")
    for step in steps:
        label = "flagged" if step == "titles" else "fixed"
        print(f"  {step:<7} {counts[step]:>4} {label}")

    fixes = sum(counts[s] for s in steps if s != "titles")

    if args.dry_run:
        print(f"\ndry run -- {fixes} change(s) not written")
        return 0

    out = args.output or f"{os.path.splitext(args.pptx)[0]}-normalized.pptx"
    prs.save(out)
    print(f"\nwrote {out} ({fixes} change(s))")
    if counts.get("titles"):
        print(f"note: {counts['titles']} title(s) need a text edit -- see TITLES above")
    return 0


CANVAS_W = T.CANVAS["width_in"]
CANVAS_H = T.CANVAS["height_in"]


if __name__ == "__main__":
    sys.exit(main())
